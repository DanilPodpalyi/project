"""Тесты чтения TXT, DOCX и PDF и схемы резюме."""

"""Тесты чтения документов и схемы резюме."""

import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from resume_evaluator.resume_extractor import (
    ResumeData,
    _is_candidate_name_placeholder,
    _reconcile_explicit_skills,
    extract_resume,
    read_resume,
    save_result,
)


class ReadResumeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.directory = Path(self.temporary_directory.name)

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_reads_utf8_txt_with_bom(self) -> None:
        file_path = self.directory / "resume.TXT"
        file_path.write_text("Имя: Иван", encoding="utf-8-sig")

        self.assertEqual(read_resume(file_path), "Имя: Иван")

    def test_reads_docx_paragraphs_and_tables_in_document_order(self) -> None:
        file_path = self.directory / "resume.docx"
        document = Document()
        document.add_paragraph("Имя: Иван")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Навык"
        table.cell(0, 1).text = "Python"
        document.add_paragraph("Опыт: 3 года")
        document.save(file_path)

        self.assertEqual(
            read_resume(file_path),
            "Имя: Иван\nНавык\tPython\nОпыт: 3 года",
        )

    def test_reads_docx_without_corrupting_unicode_or_numbers(self) -> None:
        file_path = self.directory / "unicode_resume.docx"
        document = Document()
        document.add_paragraph(
            "Проверил работу на тестовом наборе журналов "
            "объёмом более 500 000 записей."
        )
        document.add_paragraph(
            "Сократил время первичного анализа примерно на 25% в 2024 году."
        )
        document.save(file_path)

        self.assertEqual(
            read_resume(file_path),
            "Проверил работу на тестовом наборе журналов "
            "объёмом более 500 000 записей.\n"
            "Сократил время первичного анализа примерно на 25% в 2024 году.",
        )

    def test_extract_resume_preserves_unicode_and_numbers_through_json(self) -> None:
        source_result = (
            "Проверил работу на тестовом наборе журналов "
            "объёмом более 500 000 записей."
        )
        source_achievement = (
            "Сократил время первичного анализа примерно на 25%."
        )
        resume_text = f"ПРОЕКТЫ\n{source_result}\nОПЫТ\n{source_achievement}"
        payload = self._resume_data()
        payload.update({
            "target_profession": {
                "role": "Analyst",
                "specialization": "Information Security Analyst",
                "domain": "Information Security",
                "seniority": "Middle",
                "core_skills": ["SIEM"],
            },
            "work_experience": [{
                "company": "Example",
                "position": "Information Security Analyst",
                "period": "02.2024 — н.в.",
                "responsibilities": [],
                "technologies": ["SIEM"],
                "achievements": [source_achievement],
            }],
            "education": [{
                "institution": "Уральский технический университет",
                "education_level": "Бакалавриат",
                "education_field": "Информационные системы",
                "field_of_study": "Информационные системы и технологии",
                "graduation_year": 2022,
            }],
            "projects": [{
                "name": "Анализ журналов авторизации",
                "description": "Обработка журналов SSH",
                "role": "Разработчик",
                "technologies": ["Python", "Linux"],
                "results": [source_result],
            }],
        })
        response = SimpleNamespace(
            message=SimpleNamespace(
                content=json.dumps(payload, ensure_ascii=False),
            )
        )

        with patch("resume_evaluator.resume_extractor.chat", return_value=response) as chat_mock:
            extracted = extract_resume(resume_text)

        prompt = chat_mock.call_args.kwargs["messages"][-1]["content"]
        self.assertIn(source_result, prompt)
        self.assertIn(source_achievement, prompt)
        self.assertEqual(extracted.projects[0].results, [source_result])
        self.assertEqual(
            extracted.work_experience[0].achievements,
            [source_achievement],
        )
        self.assertEqual(extracted.education[0].graduation_year, 2022)

        output_path = self.directory / "resume.json"
        save_result(extracted, output_path)
        saved = ResumeData.model_validate_json(
            output_path.read_text(encoding="utf-8")
        )
        self.assertEqual(saved.projects[0].results, [source_result])
        self.assertEqual(saved.education[0].graduation_year, 2022)

    def test_reads_pdf_text_layer(self) -> None:
        file_path = self.directory / "resume.pdf"
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_reference = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 72 720 Td (Python Developer) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        with file_path.open("wb") as output:
            writer.write(output)

        self.assertEqual(read_resume(file_path), "Python Developer")

    def test_reports_pdf_without_text_layer(self) -> None:
        file_path = self.directory / "scan.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with file_path.open("wb") as output:
            writer.write(output)

        with self.assertRaisesRegex(SystemExit, "OCR"):
            read_resume(file_path)

    def test_reports_old_doc_format(self) -> None:
        file_path = self.directory / "resume.doc"
        file_path.write_bytes(b"legacy doc")

        with self.assertRaisesRegex(SystemExit, "сохранить как .docx"):
            read_resume(file_path)

    def test_reports_corrupted_docx(self) -> None:
        file_path = self.directory / "broken.docx"
        file_path.write_bytes(b"not a docx")

        with self.assertRaisesRegex(SystemExit, "Некорректный DOCX-файл"):
            read_resume(file_path)

    def test_all_schema_fields_are_required_for_model_output(self) -> None:
        schema = ResumeData.model_json_schema()

        self.assertEqual(set(schema["required"]), set(schema["properties"]))
        for definition in schema["$defs"].values():
            self.assertEqual(
                set(definition["required"]),
                set(definition["properties"]),
            )

    def test_uses_structured_target_profession(self) -> None:
        data = self._resume_data()
        data["target_profession"] = {
            "role": "Analyst",
            "specialization": "Traffic Analyst",
            "domain": "Digital Marketing / Web Analytics",
            "seniority": "Junior/Middle",
            "core_skills": ["Google Analytics", "SQL"],
        }

        resume = ResumeData.model_validate(data)

        self.assertEqual(resume.target_profession.role, "Analyst")
        self.assertEqual(resume.target_position, "Traffic Analyst")
        self.assertNotIn("target_position", resume.model_dump())

    def test_accepts_legacy_target_position(self) -> None:
        data = self._resume_data()
        data["target_position"] = "Python-разработчик"
        data["skills"] = ["Python"]

        resume = ResumeData.model_validate(data)

        self.assertEqual(
            resume.target_profession.specialization,
            "Python-разработчик",
        )
        self.assertEqual(resume.target_profession.core_skills, ["Python"])

    @staticmethod
    def _resume_data() -> dict:
        return {
            "candidate_name": "Иван",
            "professional_summary": None,
            "contact_channels": [],
            "skills": [],
            "soft_skills": [],
            "work_experience": [],
            "education": [],
            "projects": [],
            "courses_and_certifications": [],
            "languages": [],
            "achievements": [],
            "portfolio_links": [],
        }

    def test_recognizes_candidate_name_placeholders(self) -> None:
        self.assertTrue(_is_candidate_name_placeholder("ФИО"))
        self.assertTrue(
            _is_candidate_name_placeholder("Фамилия, имя, отчество")
        )
        self.assertFalse(_is_candidate_name_placeholder("Иван Петров"))

    def test_reconciles_explicit_technologies_into_skills(self) -> None:
        data = self._resume_data()
        data.update({
            "target_profession": {
                "role": "Analyst",
                "specialization": "Information Security Analyst",
                "domain": "Information Security",
                "seniority": "Middle",
                "core_skills": [],
            },
            "skills": ["Python", "Wazuh"],
            "work_experience": [{
                "company": "Example",
                "position": "Security Analyst",
                "period": None,
                "responsibilities": [],
                "technologies": ["Python", "Wireshark"],
                "achievements": [],
            }],
            "projects": [{
                "name": "SOC lab",
                "description": None,
                "role": None,
                "technologies": ["Wazuh", "VirtualBox", "Sysmon"],
                "results": [],
            }],
        })

        resume = _reconcile_explicit_skills(ResumeData.model_validate(data))

        self.assertEqual(
            resume.skills,
            ["Python", "Wazuh", "Wireshark", "VirtualBox", "Sysmon"],
        )
        self.assertEqual(resume.education, [])
        self.assertEqual(len(resume.projects), 1)


if __name__ == "__main__":
    unittest.main()
